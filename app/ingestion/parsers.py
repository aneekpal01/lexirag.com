"""Document parser implementations for PDF, DOCX, and TXT files."""

import hashlib
import io
import re
from typing import Optional
import docx
import pypdf

from app.core.exceptions import EmptyDocumentError, InvalidDocumentError
from app.core.logging import get_logger
from app.ingestion.models import DocumentPage, ParsedDocument

logger = get_logger(__name__)


def compute_sha256(content: bytes) -> str:
    """Generates canonical SHA-256 hex digest for document deduplication."""
    return hashlib.sha256(content).hexdigest()


class PDFDocumentParser:
    """Extracts text page-by-page from PDF documents using pypdf."""

    @staticmethod
    def parse(file_bytes: bytes, filename: str, domain: Optional[str] = None) -> ParsedDocument:
        stream = io.BytesIO(file_bytes)
        try:
            reader = pypdf.PdfReader(stream)
        except Exception as exc:
            logger.warning("Failed to open PDF %s: %s", filename, exc)
            raise InvalidDocumentError(f"PDF file is corrupted or unreadable: {exc}") from exc

        if reader.is_encrypted:
            try:
                # Attempt blank password decrypt
                decrypted = reader.decrypt("")
                if not decrypted:
                    raise InvalidDocumentError("PDF document is password protected and cannot be ingested.")
            except Exception as exc:
                raise InvalidDocumentError("PDF document is encrypted or password protected.") from exc

        pages: list[DocumentPage] = []
        full_text_parts: list[str] = []

        try:
            for page_index, page in enumerate(reader.pages, start=1):
                extracted = page.extract_text() or ""
                cleaned = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f]", "", extracted).strip()
                if cleaned:
                    pages.append(DocumentPage(page_number=page_index, text=cleaned))
                    full_text_parts.append(cleaned)
        except Exception as exc:
            logger.warning("Error reading pages from PDF %s: %s", filename, exc)
            raise InvalidDocumentError(f"Error while extracting text from PDF: {exc}") from exc

        concatenated_text = "\n\n".join(full_text_parts).strip()
        if not concatenated_text:
            raise EmptyDocumentError(
                f"PDF '{filename}' yielded zero extractable characters. Scanned image PDFs without OCR are not supported."
            )

        return ParsedDocument(
            document_id=compute_sha256(file_bytes),
            filename=filename,
            file_type="pdf",
            file_size_bytes=len(file_bytes),
            pages=pages,
            domain=domain,
            raw_text=concatenated_text,
        )


class DocxDocumentParser:
    """Extracts structured text from Microsoft Word .docx files."""

    @staticmethod
    def parse(file_bytes: bytes, filename: str, domain: Optional[str] = None) -> ParsedDocument:
        stream = io.BytesIO(file_bytes)
        try:
            document = docx.Document(stream)
        except Exception as exc:
            logger.warning("Failed to open DOCX %s: %s", filename, exc)
            raise InvalidDocumentError(f"DOCX document is corrupted or not a valid Word archive: {exc}") from exc

        extracted_lines: list[str] = []

        # 1. Paragraphs
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            # Preserve heading structure notation if defined
            if para.style and para.style.name and para.style.name.lower().startswith("heading"):
                extracted_lines.append(f"\n## {text}\n")
            else:
                extracted_lines.append(text)

        # 2. Table cells
        for table in document.tables:
            table_lines: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_lines.append(" | ".join(row_cells))
            if table_lines:
                extracted_lines.append("\n[TABLE DATA]\n" + "\n".join(table_lines) + "\n")

        concatenated_text = "\n\n".join(extracted_lines).strip()
        if not concatenated_text:
            raise EmptyDocumentError(f"DOCX file '{filename}' contains no readable text.")

        # In DOCX, physical page boundaries are dynamic; we treat as segment 1
        pages = [DocumentPage(page_number=1, text=concatenated_text)]

        return ParsedDocument(
            document_id=compute_sha256(file_bytes),
            filename=filename,
            file_type="docx",
            file_size_bytes=len(file_bytes),
            pages=pages,
            domain=domain,
            raw_text=concatenated_text,
        )


class PlainTextParser:
    """Parses plain UTF-8 or Latin-1 text files."""

    @staticmethod
    def parse(file_bytes: bytes, filename: str, domain: Optional[str] = None) -> ParsedDocument:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1")
            except Exception as exc:
                raise InvalidDocumentError(f"Failed to decode text file '{filename}': {exc}") from exc

        cleaned_text = text.strip()
        if not cleaned_text:
            raise EmptyDocumentError(f"Text file '{filename}' is empty.")

        pages = [DocumentPage(page_number=1, text=cleaned_text)]

        return ParsedDocument(
            document_id=compute_sha256(file_bytes),
            filename=filename,
            file_type="txt",
            file_size_bytes=len(file_bytes),
            pages=pages,
            domain=domain,
            raw_text=cleaned_text,
        )


def parse_document(file_bytes: bytes, filename: str, domain: Optional[str] = None) -> ParsedDocument:
    """Dispatches raw document bytes to the appropriate parser based on file extension."""
    if not file_bytes:
        raise EmptyDocumentError(f"Document payload for '{filename}' is completely empty (0 bytes).")

    lowered_name = filename.lower().strip()
    if lowered_name.endswith(".pdf"):
        return PDFDocumentParser.parse(file_bytes, filename, domain)
    elif lowered_name.endswith(".docx"):
        return DocxDocumentParser.parse(file_bytes, filename, domain)
    elif lowered_name.endswith(".txt"):
        return PlainTextParser.parse(file_bytes, filename, domain)
    else:
        raise InvalidDocumentError(
            f"Unsupported file format for '{filename}'. Allowed formats: .pdf, .docx, .txt"
        )
